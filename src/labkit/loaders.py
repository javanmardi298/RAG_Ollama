"""Read files into Documents (text + metadata). PDF and DOCX readers are optional extras.

Real documents are messy: headers and footers repeat on every page, hyphenated line breaks split words,
tables lose their layout. `clean_text` fixes the cheap cases; check a sample by eye before indexing.
"""
from __future__ import annotations

import html
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Document:
    doc_id: str
    text: str
    meta: dict = field(default_factory=dict)      # source, page, title, lang, ...


_TAG = re.compile(r"<[^>]+>")
_SCRIPT = re.compile(r"<(script|style)[^>]*>.*?</\1>", re.S | re.I)
_HEADING = re.compile(r"<h([1-6])[^>]*>(.*?)</h\1>", re.S | re.I)


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("­", "")       # soft hyphens
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)                   # words split at line end
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def drop_repeated_lines(pages: list[str], min_share: float = 0.6) -> list[str]:
    """Remove lines that appear on most pages (running headers, footers, page numbers)."""
    if len(pages) < 3:
        return pages
    counts = Counter(line.strip() for p in pages for line in set(p.splitlines()) if line.strip())
    repeated = {l for l, c in counts.items() if c / len(pages) >= min_share}
    number = re.compile(r"^(page\s*)?\d+(\s*/\s*\d+)?$", re.I)
    return ["\n".join(l for l in p.splitlines() if l.strip() not in repeated and not number.match(l.strip()))
            for p in pages]


def html_to_text(raw: str) -> str:
    """Keep headings as markdown '#' lines so heading-based chunking still works."""
    raw = _SCRIPT.sub(" ", raw)
    raw = _HEADING.sub(lambda m: "\n\n" + "#" * int(m.group(1)) + " " + _TAG.sub("", m.group(2)).strip() + "\n\n", raw)
    raw = re.sub(r"<(br|/p|/div|/li|/tr)[^>]*>", "\n", raw, flags=re.I)
    return clean_text(html.unescape(_TAG.sub(" ", raw)))


def load_file(path: str | Path) -> list[Document]:
    """One Document per file (per page for PDF, so citations can point to a page)."""
    path = Path(path)
    ext = path.suffix.lower()
    base = {"source": str(path), "title": path.stem}
    if ext in {".txt", ".md", ".markdown", ".rst"}:
        return [Document(path.stem, clean_text(path.read_text(encoding="utf-8", errors="replace")), base)]
    if ext in {".html", ".htm"}:
        return [Document(path.stem, html_to_text(path.read_text(encoding="utf-8", errors="replace")), base)]
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise ImportError("PDF needs: pip install pypdf") from e
        pages = drop_repeated_lines([p.extract_text() or "" for p in PdfReader(path).pages])
        return [Document(f"{path.stem}#p{i}", clean_text(t), {**base, "page": i})
                for i, t in enumerate(pages, 1) if t.strip()]
    if ext == ".docx":
        try:
            import docx
        except ImportError as e:
            raise ImportError("DOCX needs: pip install python-docx") from e
        d = docx.Document(path)
        lines = [("#" * int(p.style.name[-1]) + " " if p.style.name.startswith("Heading ") and p.style.name[-1].isdigit()
                  else "") + p.text for p in d.paragraphs]
        return [Document(path.stem, clean_text("\n\n".join(lines)), base)]
    raise ValueError(f"unsupported file type: {ext}")


def load_dir(folder: str | Path, patterns=("*.md", "*.txt", "*.html", "*.pdf", "*.docx")) -> list[Document]:
    docs = []
    for pattern in patterns:
        for p in sorted(Path(folder).rglob(pattern)):
            docs += load_file(p)
    return docs
